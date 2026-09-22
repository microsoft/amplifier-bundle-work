"""Native feature preservation for original fixtures; no visual claims."""
import importlib.util
import json
from pathlib import Path
import pytest

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('fidelity', ROOT/'evals/fidelity.py')
fidelity=importlib.util.module_from_spec(spec);spec.loader.exec_module(fidelity)

@pytest.fixture
def fixtures(tmp_path):
    for package in ('docx','pptx','openpyxl'): pytest.importorskip(package)
    fidelity.prepare(tmp_path)
    return tmp_path


def test_styled_document_edit_preserves_native_features(fixtures):
    from docx import Document
    source=fixtures/'sources/styled.docx';before=fidelity.digest(source)
    doc=Document(source)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if 'Cedar' in run.text:run.text=run.text.replace('Cedar','Aspen')
    output=fixtures/'edited.docx';doc.save(output)
    assert fidelity.inventory(output)==fidelity.inventory(source)
    assert fidelity.digest(source)==before
    assert 'Aspen' in '\n'.join(p.text for p in Document(output).paragraphs)


def test_chart_and_notes_remain_editable(fixtures):
    from pptx import Presentation
    source=fixtures/'sources/editable.pptx';before=fidelity.digest(source)
    deck=Presentation(source);deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text='Aspen evidence'
    output=fixtures/'edited.pptx';deck.save(output)
    assert fidelity.inventory(output)==fidelity.inventory(source)
    assert fidelity.digest(source)==before


def test_workbook_input_preserves_names_rules_hidden_sheet_chart(fixtures):
    from openpyxl import load_workbook
    source=fixtures/'sources/model.xlsx';before=fidelity.digest(source)
    book=load_workbook(source);book['Inputs']['F2']=.15
    output=fixtures/'edited.xlsx';book.save(output)
    assert fidelity.inventory(output)==fidelity.inventory(source)
    assert fidelity.digest(source)==before
    assert load_workbook(output)['Inputs']['F2'].value==.15


def test_fixtures_refuse_replacement(fixtures):
    with pytest.raises(FileExistsError):fidelity.prepare(fixtures)


def test_natural_cases_include_ordinary_turns_and_near_misses():
    cases=json.loads((ROOT/'evals/natural-cases.json').read_text())['cases']
    assert len({c['id'] for c in cases})==len(cases)
    assert len([c for c in cases if c['kind']=='sustained' and len(c['turns'])==2])==3
    assert len([c for c in cases if c['kind']=='near-miss' and c['expected_skills']==[]])>=2
    for case in cases:
        assert 'load_skill' not in ' '.join(case['turns'])
        assert all(name in {p.parent.name for p in (ROOT/'skills').glob('*/SKILL.md')} for name in case['expected_skills'])


def test_changed_named_input_recalculates_in_real_engine(fixtures):
    import os, shutil, subprocess
    from openpyxl import load_workbook
    office=os.environ.get('WORK_SOFFICE') or shutil.which('soffice')
    if not office:pytest.skip('Actual spreadsheet engine required')
    source=fixtures/'sources/model.xlsx';before=fidelity.digest(source)
    for rate,expected in [(.1,58.5),(.15,55.25),(0,65.0)]:
        folder=fixtures/str(rate);folder.mkdir();dest=folder/'calculated';dest.mkdir()
        book=load_workbook(source);book['Inputs']['F2']=rate;draft=folder/'model.xlsx';book.save(draft)
        subprocess.run([office,f'-env:UserInstallation={(folder/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(dest),str(draft)],check=True,capture_output=True,timeout=90)
        calculated=load_workbook(dest/'model.xlsx',data_only=True)
        assert calculated['Inputs']['F5'].value==pytest.approx(expected)
        assert calculated['Inputs']['D4'].value==-12.5
        assert calculated['Inputs']['D5'].value==0
        formulas=load_workbook(dest/'model.xlsx',data_only=False)
        assert formulas['Inputs']['F5'].data_type=='f'
        assert formulas['Notes'].sheet_state=='hidden'
        assert len(formulas['Inputs']._charts)==1
    assert fidelity.digest(source)==before


@pytest.mark.parametrize('numbers,duration,review,failed', [
    ('150 tickets, 115 second lookups', 'two-week pilot', '12 October 2026', {'sample','lookups'}),
    ('50 tickets, 15 second lookups', 'three-week pilot', '19 October 2026', {'duration','review'}),
    ('50 tickets, 15 second lookups', 'twenty-two-week pilot', '112 October 2026', {'duration','review'}),
])
def test_memo_oracle_rejects_wrong_numbers_and_dates(fixtures, numbers, duration, review, failed):
    from docx import Document
    output=fixtures/'outputs';output.mkdir()
    doc=Document();doc.add_heading('Decision',0)
    doc.add_paragraph(f'Approve {duration}. Mara Singh. Evidence: {numbers}. Review: {review}. Source: evidence-current.csv.')
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Missing owner';table.cell(0,1).text='Require an owner field'
    doc.save(output/'decision.docx')
    result=fidelity.verify_memo(fixtures)
    assert all(result[name] is False for name in failed)
    assert result['sources_unchanged']
