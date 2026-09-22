"""Real engine results in delivered files, preserving source OOXML features."""
import importlib.util
import os
from pathlib import Path
import shutil
import zipfile

import pytest

ROOT = Path(__file__).resolve().parents[1]


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, ROOT/path)
    value = importlib.util.module_from_spec(spec); spec.loader.exec_module(value)
    return value


finalizer = module('xlsx_finalizer', 'skills/spreadsheets/scripts/finalize_xlsx.py')
fidelity = module('xlsx_fixture', 'evals/fidelity.py')


@pytest.fixture
def calculated(tmp_path):
    from openpyxl import load_workbook
    if not (os.environ.get('WORK_SOFFICE') or shutil.which('soffice')):
        pytest.skip('Actual spreadsheet engine required')
    fidelity.prepare(tmp_path)
    source = tmp_path/'draft.xlsx'
    book = load_workbook(tmp_path/'sources/model.xlsx')
    book['Inputs']['F2'] = .15
    book['Inputs']['B6'] = None
    book['Inputs']['D6'] = '=IF(B6="","",B6*C6)'
    book['Inputs']._charts[0].series[0].invertIfNegative = False
    book.save(source)
    result = tmp_path/'final.xlsx'
    receipt = finalizer.finalize(source, result)
    return source, result, receipt


def test_final_file_contains_actual_results_and_retains_native_parts(calculated):
    from openpyxl import load_workbook
    source, result, receipt = calculated
    cached, formulas = load_workbook(result, data_only=True), load_workbook(result)
    assert cached['Inputs']['F5'].value == pytest.approx(55.25)
    assert cached['Inputs']['D4'].value == -12.5
    assert cached['Inputs']['D5'].value == 0
    assert cached['Inputs']['D6'].value is None  # Real engine empty-string result, not zero.
    assert formulas['Inputs']['F5'].value == '=SUM(D2:D5)*(1-Discount)'
    assert fidelity.inventory(result) == fidelity.inventory(source)
    assert receipt['formulaCaches'] == 6 and receipt['calculationEvidence'].startswith('LibreOffice')
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result) as after:
        assert before.namelist() == after.namelist()
        changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
        assert changed == set(receipt['changedParts']) == {'xl/worksheets/sheet1.xml'}
        assert b'invertIfNegative val="0"' in after.read('xl/charts/chart1.xml')
    assert finalizer.digest(source) == receipt['sourceSha256']


@pytest.mark.parametrize('coordinate,value', [('F2', .2), ('F5', '=SUM(D2:D5)')])
def test_rejects_stale_input_and_formula_caches(calculated, coordinate, value):
    from openpyxl import load_workbook
    source, result, _ = calculated
    changed = source.with_name('changed.xlsx')
    book = load_workbook(source); book['Inputs'][coordinate] = value; book.save(changed)
    output = source.with_name('invalid.xlsx')
    with pytest.raises(ValueError, match='inputs or formulas differ'):
        finalizer.merge_caches(changed, result, output)
    assert not output.exists()


def test_missing_caches_and_replacement_are_refused(tmp_path):
    from openpyxl import Workbook
    source = tmp_path/'draft.xlsx'; book = Workbook(); book.active['A1'] = '=1+1'; book.save(source)
    with pytest.raises(ValueError, match='finite engine result'):
        finalizer.merge_caches(source, source, tmp_path/'final.xlsx')
    assert not (tmp_path/'final.xlsx').exists()
    with pytest.raises(FileExistsError):
        finalizer.merge_caches(source, source, source)


@pytest.mark.parametrize('dependency', ['name', 'table-range', 'table-column'])
def test_rejects_stale_calculation_dependencies(calculated, dependency):
    from openpyxl import load_workbook
    source, result, _ = calculated
    changed = source.with_name('changed-dependency.xlsx')
    book = load_workbook(source)
    if dependency == 'name':
        book.defined_names['Discount'].attr_text = "'Inputs'!$F$3"
    elif dependency == 'table-range':
        book['Inputs'].tables['Orders'].ref = 'A1:D6'
    else:
        book['Inputs'].tables['Orders'].tableColumns[-1].name = 'Wrong Revenue'
    book.save(changed)
    output = source.with_name('invalid-dependency.xlsx')
    with pytest.raises(ValueError, match='defined names differ|table dependencies differ'):
        finalizer.merge_caches(changed, result, output)
    assert not output.exists()


@pytest.mark.parametrize('change', ['input', 'format'])
def test_source_mutation_during_engine_calculation_is_refused(tmp_path, monkeypatch, change):
    from openpyxl import Workbook
    source = tmp_path/'source.xlsx'; book = Workbook(); book.active['A1'] = 1; book.save(source)
    monkeypatch.setenv('WORK_SOFFICE', __file__)
    def mutate(*args, **kwargs):
        if change == 'input':
            book.active['A1'] = 2
        else:
            book.active['A1'].number_format = '0.000'
        book.save(source)
    monkeypatch.setattr(finalizer.subprocess, 'run', mutate)
    with pytest.raises(ValueError, match='Source changed during recalculation'):
        finalizer.finalize(source, tmp_path/'final.xlsx')
    assert not (tmp_path/'final.xlsx').exists()
