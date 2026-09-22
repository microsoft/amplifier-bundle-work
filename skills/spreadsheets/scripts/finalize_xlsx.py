#!/usr/bin/env python3
"""Deliver real formula caches without round-tripping the original OOXML features."""
import argparse
import copy
import hashlib
import io
import json
import math
import os
from pathlib import Path
import posixpath
import shutil
import subprocess
import tempfile
from xml.dom.minidom import parseString
import zipfile

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def xml(data):
    if b'<!DOCTYPE' in data.upper():
        raise ValueError('OOXML document types are unsupported')
    return parseString(data)


def children(node, name):
    return [n for n in node.childNodes if n.nodeType == n.ELEMENT_NODE
            and n.namespaceURI == NS and n.localName == name]


def text(node):
    return ''.join(n.data for n in node.childNodes if n.nodeType in (n.TEXT_NODE, n.CDATA_SECTION_NODE))


def sheets(package):
    relations = xml(package.read('xl/_rels/workbook.xml.rels'))
    targets = {n.getAttribute('Id'): n.getAttribute('Target') for n in relations.documentElement.childNodes
               if n.nodeType == n.ELEMENT_NODE and n.getAttribute('TargetMode') != 'External'}
    result = {}
    for node in xml(package.read('xl/workbook.xml')).getElementsByTagNameNS(NS, 'sheet'):
        target = targets[node.getAttributeNS(REL, 'id')]
        result[node.getAttribute('name')] = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/'+target)
    return result


def aligned(source, calculated):
    from openpyxl import load_workbook
    original, engine = load_workbook(source), load_workbook(calculated)
    try:
        if original.sheetnames != engine.sheetnames:
            raise ValueError('Calculated sheet identities or order changed')
        if original.epoch != engine.epoch:
            raise ValueError('Calculated date system differs')
        def names(book):
            # Print areas/titles are parsed separately by openpyxl and do not
            # change cell calculation. Keep calculation names scoped by sheet.
            def target(item):
                if item.type == 'RANGE':
                    try:
                        # Quoted versus unquoted sheet names are equivalent;
                        # retain reference order and absolute/relative markers.
                        destinations = tuple(item.destinations)
                        if destinations:
                            return ('ranges', destinations)
                    except (AttributeError, TypeError, ValueError):
                        pass
                return ('expression', item.attr_text)
            result = {('workbook', name): target(item) for name, item in book.defined_names.items()}
            for sheet in book:
                result.update({(sheet.title, name): target(item) for name, item in sheet.defined_names.items()})
            return result
        if names(original) != names(engine):
            raise ValueError('Calculated defined names differ; refuse stale caches')
        def tables(sheet):
            def formula(value):
                return None if value is None else (value.attr_text, bool(value.array))
            return sorted((table.name, table.displayName, table.ref, table.headerRowCount,
                           table.totalsRowCount or 0,
                           tuple((column.name, column.totalsRowFunction,
                                  formula(column.calculatedColumnFormula), formula(column.totalsRowFormula))
                                 for column in table.tableColumns))
                          for table in sheet.tables.values())
        for sheet in original:
            if tables(sheet) != tables(engine[sheet.title]):
                raise ValueError(f'Calculated table dependencies differ in {sheet.title}; refuse stale caches')
            before = {c.coordinate: c.value for row in sheet for c in row if c.value is not None}
            after = {c.coordinate: c.value for row in engine[sheet.title] for c in row if c.value is not None}
            if before != after:
                raise ValueError(f'Calculated inputs or formulas differ in {sheet.title}; refuse stale caches')
    finally:
        original.close(); engine.close()


def merge_caches(source, calculated, output):
    """Copy only verified <v>/type fields; preserve every other package part."""
    source, calculated, output = Path(source).resolve(), Path(calculated).resolve(), Path(output).resolve()
    if any(p.suffix.lower() != '.xlsx' for p in (source, calculated, output)):
        raise ValueError('Only ordinary XLSX files are supported')
    if output.exists() or output in (source, calculated):
        raise FileExistsError('Use a new final output path; originals are preserved')
    source_bytes, calculated_bytes = source.read_bytes(), calculated.read_bytes()
    aligned(io.BytesIO(source_bytes), io.BytesIO(calculated_bytes))
    count, changed = 0, {}
    with zipfile.ZipFile(io.BytesIO(source_bytes)) as original, zipfile.ZipFile(io.BytesIO(calculated_bytes)) as engine:
        before, after = sheets(original), sheets(engine)
        if list(before) != list(after):
            raise ValueError('Calculated sheet identities changed')
        for name, part in before.items():
            document = xml(original.read(part))
            other = xml(engine.read(after[name]))
            lookup = {n.getAttribute('r'): n for n in other.getElementsByTagNameNS(NS, 'c')}
            updates = 0
            for cell in document.getElementsByTagNameNS(NS, 'c'):
                formula = children(cell, 'f')
                if not formula:
                    continue
                if formula[0].getAttribute('t') in {'array', 'shared', 'dataTable'}:
                    raise ValueError('Array, shared and data-table formula caches require their native engine')
                computed = lookup.get(cell.getAttribute('r'))
                if computed is None or len(children(computed, 'f')) != 1 or text(formula[0]) != text(children(computed, 'f')[0]):
                    raise ValueError(f'Calculated formula does not align: {name}!{cell.getAttribute("r")}')
                values = children(computed, 'v')
                kind = computed.getAttribute('t') or 'n'
                if len(values) != 1 or kind not in {'n', 'b', 'str'}:
                    raise ValueError(f'Missing, unsupported or error cache: {name}!{cell.getAttribute("r")}')
                value = text(values[0])
                if kind == 'n' and (not value or not math.isfinite(float(value))):
                    raise ValueError('Numeric formula cache must contain a finite engine result')
                if kind == 'b' and value not in {'0', '1'}:
                    raise ValueError('Boolean formula cache is invalid')
                for old in children(cell, 'v') + children(cell, 'is'):
                    cell.removeChild(old)
                # Use the original worksheet prefix/namespace, even if the
                # engine serialized its worksheet with different XML prefixes.
                cached = document.createElementNS(NS, (cell.prefix+':' if cell.prefix else '')+'v')
                if value:
                    cached.appendChild(document.createTextNode(value))
                cell.insertBefore(cached, formula[0].nextSibling)
                cell.setAttribute('t', kind)
                count += 1; updates += 1
            if updates:
                changed[part] = document.toxml(encoding='UTF-8')
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.xlsx') as staging:
            with zipfile.ZipFile(staging.name, 'w') as target:
                for info in original.infolist():
                    target.writestr(copy.copy(info), changed.get(info.filename, original.read(info.filename)))
            # Exclusive publication: an existing output is never replaced.
            os.link(staging.name, output)
    return {'output': str(output), 'formulaCaches': count, 'changedParts': sorted(changed),
            'sourceSha256': hashlib.sha256(source_bytes).hexdigest(), 'calculatedSha256': hashlib.sha256(calculated_bytes).hexdigest(),
            'outputSha256': digest(output), 'visualInspection': 'required for every final page'}


def finalize(source, output, calculated=None):
    if calculated:
        return {**merge_caches(source, calculated, output), 'calculationEvidence': 'provided engine output; caller verifies provenance'}
    office = os.environ.get('WORK_SOFFICE') or shutil.which('soffice')
    if not office or not Path(office).is_file():
        raise RuntimeError('Install LibreOffice or set WORK_SOFFICE to its executable')
    with tempfile.TemporaryDirectory(prefix='work-calculate-') as temporary:
        root = Path(temporary); inputs = root/'input'; results = root/'calculated'
        inputs.mkdir(); results.mkdir()
        draft = inputs/'workbook.xlsx'; source_bytes = Path(source).read_bytes(); draft.write_bytes(source_bytes)
        subprocess.run([office, f'-env:UserInstallation={(root/"profile").as_uri()}', '--headless',
                        '--convert-to', 'xlsx', '--outdir', str(results), str(draft)],
                       check=True, capture_output=True, text=True, timeout=120)
        if Path(source).read_bytes() != source_bytes:
            raise ValueError('Source changed during recalculation; refuse stale caches')
        # Publish from the exact snapshot handed to the engine. A source edit
        # after this check cannot introduce uncalculated inputs or metadata.
        return {**merge_caches(draft, results/'workbook.xlsx', output), 'calculationEvidence': 'LibreOffice completed in isolated profile'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source'); parser.add_argument('--output', required=True)
    parser.add_argument('--calculated', help='Previously observed real engine output for the same inputs/formulas')
    args = parser.parse_args()
    try:
        print(json.dumps(finalize(args.source, args.output, args.calculated), indent=2))
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError, zipfile.BadZipFile) as error:
        parser.exit(1, f'Workbook finalization failed: {error}\n')
