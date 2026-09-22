"""Original portable fidelity fixtures; no model calls or proprietary assets."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def prepare(destination):
    """Materialize original inputs once; refuse overwrites."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    root = Path(destination) / 'sources'
    root.mkdir(parents=True, exist_ok=False)
    (root / 'brief.md').write_text('# Cedar launch decision\nOwner: Mara Singh\nAudience: Operations\nDecision: approve a two-week pilot.\nStart: 28 September 2026\nReview: 12 October 2026\nUse evidence-current.csv; evidence-old.csv is superseded.\nRisk: Missing owner. Mitigation: require an owner field.\nRisk: Stale status. Mitigation: display a freshness timestamp.\nAll facts are synthetic.\n')
    (root / 'evidence-current.csv').write_text('metric,value\ninspected tickets,50\nsecond lookups,15\n')
    (root / 'evidence-old.csv').write_text('metric,value\ninspected tickets,40\nsecond lookups,24\n')
    (root / 'metrics.csv').write_text('week,tickets,second_lookups\nWeek 1,20,8\nWeek 2,30,7\n')
    (root / 'orders.csv').write_text('date,item,quantity,unit_price\n2026-09-01,Alpha,3,12.50\n2026-09-02,Beta,2,20.00\n2026-09-03,Refund,-1,12.50\n2026-09-04,Zero,0,8.00\n2026-09-05,Pending,,9.00\n')
    (root / 'correction.csv').write_text('date,item,quantity,unit_price\n2026-09-06,Gamma,4,5.00\n')
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'Operations | Internal'
    doc.sections[0].footer.paragraphs[0].text = 'Synthetic fidelity fixture'
    doc.add_heading('Cedar decision', 0)
    paragraph = doc.add_paragraph()
    paragraph.add_run('Cedar').bold = True
    paragraph.add_run(' has a named owner and a dated review.').italic = True
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), '1'); start.set(qn('w:name'), 'decision')
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), '1')
    paragraph._p.insert(0, start); paragraph._p.append(end)
    rel = paragraph.part.relate_to('https://example.org/cedar', RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Source brief'; run.append(text); link.append(run); paragraph._p.append(link)
    table = doc.add_table(rows=2, cols=2); table.style = 'Table Grid'
    for row, values in zip(table.rows, [('Risk', 'Mitigation'), ('Missing owner', 'Require owner field')]):
        for cell, value in zip(row.cells, values): cell.text = value
    doc.save(root / 'styled.docx')

    deck = Presentation(); deck.slide_width = Inches(12); deck.slide_height = Inches(6.75)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(10.8), Inches(.7))
    title.text_frame.paragraphs[0].font.size = Pt(30)
    title.text_frame.paragraphs[0].add_run().text = 'Cedar evidence'
    data = CategoryChartData(); data.categories = ['Week 1', 'Week 2']; data.add_series('Second lookups', [8, 7])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.8), Inches(1.5), Inches(10), Inches(4.5), data)
    slide.notes_slide.notes_text_frame.text = 'Source: metrics.csv. Synthetic values.'
    deck.save(root / 'editable.pptx')

    book = Workbook(); sheet = book.active; sheet.title = 'Inputs'
    sheet.append(['Item', 'Units', 'Unit price', 'Revenue'])
    for values in [('Alpha',3,12.5),('Beta',2,20),('Refund',-1,12.5),('Zero',0,8)]: sheet.append(list(values))
    for row in range(2,6):
        sheet.cell(row,4,f'=B{row}*C{row}')
        for column in (3,4): sheet.cell(row,column).number_format = '$#,##0.00;[Red]($#,##0.00)'
    sheet['F1'] = 'Discount'; sheet['F2'] = .1; sheet['F2'].number_format = '0%'
    sheet['F4'] = 'Net revenue'; sheet['F5'] = '=SUM(D2:D5)*(1-Discount)'
    sheet['F5'].number_format = '$#,##0.00'
    book.defined_names.add(DefinedName('Discount', attr_text="'Inputs'!$F$2"))
    validation = DataValidation(type='decimal',operator='between',formula1='0',formula2='1'); validation.add(sheet['F2']); sheet.add_data_validation(validation)
    table = Table(displayName='Orders',ref='A1:D5'); table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); sheet.add_table(table)
    sheet.conditional_formatting.add('D2:D5', CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='FFDADA')))
    chart = BarChart(); chart.add_data(Reference(sheet,min_col=4,min_row=1,max_row=5),titles_from_data=True); chart.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=5)); chart.y_axis.title='Revenue (USD)'; chart.x_axis.title='Item'; sheet.add_chart(chart,'A8')
    sheet.freeze_panes='A2'; sheet.print_area='A1:H23'; sheet.page_setup.orientation='landscape'; sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=1; sheet.sheet_properties.pageSetUpPr.fitToPage=True
    for column in 'ABCDEFGH': sheet.column_dimensions[column].width=17
    notes=book.create_sheet('Notes'); notes['A1']='Synthetic fixture; keep hidden.'; notes.sheet_state='hidden'
    book.save(root / 'model.xlsx')
    hashes={path.name:digest(path) for path in sorted(root.iterdir())}
    (Path(destination)/'source-hashes.json').write_text(json.dumps(hashes,indent=2)+'\n')
    return hashes


def inventory(path):
    """Stable semantic preservation oracle; visual/native-engine checks separate."""
    path=Path(path)
    if path.suffix == '.docx':
        from docx import Document
        doc=Document(path)
        from docx.oxml.ns import qn
        return {'headers':[s.header.paragraphs[0].text for s in doc.sections],
            'footers':[s.footer.paragraphs[0].text for s in doc.sections],
            'links':sorted(r.target_ref for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink')),
            'bookmarks':[node.get(qn('w:name')) for node in doc.element.iter(qn('w:bookmarkStart'))],
            'tables':[[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables],
            'run_styles':[(run.bold,run.italic,run.underline) for p in doc.paragraphs for run in p.runs],
            'paragraph_styles':[p.style.name for p in doc.paragraphs]}
    if path.suffix == '.pptx':
        from pptx import Presentation
        deck=Presentation(path)
        return {'slides':len(deck.slides),'size':[deck.slide_width,deck.slide_height],
            'charts':[[list(series.values) for series in shape.chart.series] for slide in deck.slides for shape in slide.shapes if shape.has_chart],
            'notes':[slide.notes_slide.notes_text_frame.text for slide in deck.slides],
            'objects':[[[shape.shape_type,shape.left,shape.top,shape.width,shape.height] for shape in slide.shapes] for slide in deck.slides]}
    if path.suffix == '.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path); sheet=book['Inputs']
        return {'sheets':[(s.title,s.sheet_state) for s in book], 'names':[(k,v.attr_text) for k,v in book.defined_names.items()],
            'validation':[(v.type,v.operator,v.formula1,v.formula2,str(v.sqref)) for v in sheet.data_validations.dataValidation],
            'tables':[(t.name,t.ref) for t in sheet.tables.values()], 'conditional_formats':len(sheet.conditional_formatting),
            'charts':len(sheet._charts), 'formula':sheet['F5'].value, 'print_area':str(sheet.print_area), 'freeze':sheet.freeze_panes}
    raise ValueError('Unsupported fidelity fixture')


def verify_memo(workspace, revised=False):
    from docx import Document
    root=Path(workspace); path=root/'outputs'/('decision-revised.docx' if revised else 'decision.docx')
    doc=Document(path)
    text=' '.join([p.text for p in doc.paragraphs]+[cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    normalized=' '.join(text.lower().split())
    checks={key:value in normalized for key,value in {'owner':'mara singh','decision':'approve','sample':'50','lookups':'15','source':'evidence-current.csv','risk':'missing owner','mitigation':'require an owner field'}.items()}
    checks['duration']=('three' in normalized or '3-week' in normalized or '3 week' in normalized) if revised else ('two' in normalized or '2-week' in normalized or '2 week' in normalized)
    checks['review']=('19 october' in normalized or 'october 19' in normalized or '2026-10-19' in normalized) if revised else ('12 october' in normalized or 'october 12' in normalized or '2026-10-12' in normalized)
    checks['table']=bool(doc.tables)
    checks['semantic_title']=any(p.style.name in {'Title','Heading 1'} for p in doc.paragraphs)
    hashes=json.loads((root/'source-hashes.json').read_text());checks['sources_unchanged']=all(digest(root/'sources'/name)==value for name,value in hashes.items())
    return checks


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('action',choices=['prepare','inventory','memo','memo-revised']);parser.add_argument('path',type=Path);args=parser.parse_args()
    result=prepare(args.path) if args.action=='prepare' else inventory(args.path) if args.action=='inventory' else verify_memo(args.path,args.action=='memo-revised')
    print(json.dumps(result,indent=2))
