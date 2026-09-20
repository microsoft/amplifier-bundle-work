#!/usr/bin/env python3
"""Build original Work reference templates. Does not read or copy OpenAI assets."""
import importlib.util
from pathlib import Path
import tempfile

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('create_template', ROOT / 'skills/template-creator/scripts/create_template.py')
creator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(creator)

DOCUMENTS = {
    'design-report': ['Problem and audience', 'Requirements', 'Design approach', 'Alternatives', 'Validation and next steps'],
    'experiment-analysis': ['Question and hypothesis', 'Method and sample', 'Results', 'Limitations', 'Decision and follow-up'],
    'investment-committee-memo': ['Proposed investment', 'Business and market', 'Financial analysis', 'Risks and mitigations', 'Recommendation'],
    'legal-memorandum': ['Question presented', 'Relevant facts', 'Applicable authorities', 'Analysis', 'Conclusion'],
    'minimal-letterhead': ['Recipient and date', 'Subject', 'Message', 'Next step and signature'],
    'strategy-memorandum': ['Decision required', 'Current position', 'Options and tradeoffs', 'Recommended approach', 'Execution and measures'],
    'system-design': ['Goals and constraints', 'System context', 'Components and interfaces', 'Data and failure handling', 'Validation and rollout'],
}
DECKS = {
    'business-review': ['Executive assessment', 'Results and evidence', 'Drivers and risks', 'Decisions and owners'],
    'market-trends-report': ['Market scope', 'Observed trends', 'Evidence and uncertainty', 'Implications'],
    'operating-review': ['Operating priorities', 'Performance', 'Issues and dependencies', 'Actions'],
    'project-kickoff': ['Purpose and outcomes', 'Scope and deliverables', 'Milestones and roles', 'Risks and next steps'],
    'simple-dark-mode': ['Main point', 'Supporting evidence', 'Recommendation'],
    'simple-light-mode': ['Main point', 'Supporting evidence', 'Recommendation'],
    'team-alignment': ['Shared objective', 'Responsibilities', 'Dependencies and decisions', 'Commitments'],
}
SHEETS = ['analytics-dashboard', 'financial-budget', 'operating-calendar',
          'project-tracker', 'sales-pipeline', 'three-statement-forecast']


def document(name, sections, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(.7)
    doc.sections[0].left_margin = doc.sections[0].right_margin = Inches(.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for style in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[style].font.name = 'Arial'
        doc.styles[style].font.color.rgb = RGBColor.from_string('183B56')
    doc.add_paragraph(name.replace('-', ' ').title(), 'Title')
    doc.add_paragraph('Author: [name]    Date: [date]    Audience: [reader]')
    for section in sections:
        doc.add_heading(section, 1)
        doc.add_paragraph(f'[Write the {section.lower()} using verified sources and the reader’s needs.]')
    doc.add_paragraph('Sources: [identify the evidence and relevant dates]')
    doc.save(path)


def deck(name, sections, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    dark = name == 'simple-dark-mode'
    foreground = 'F1F5F9' if dark else '183B56'
    background = '17212B' if dark else 'FFFFFF'
    for index, heading in enumerate([name.replace('-', ' ').title(), *sections]):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        for x,y,w,h,text,size in [
            (.8,.7,11.7,1.1,heading,32),
            (.8,2.2,11.5,2.8,'[Add the main message and supporting evidence.]',23),
            (.8,6.7,10,.35,'[Source or presenter notes]',11),
            (12,6.7,.5,.35,str(index+1),11),
        ]:
            shape = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
            shape.text_frame.word_wrap = True
            p = shape.text_frame.paragraphs[0]
            p.text = text
            p.font.name = 'Arial'
            p.font.size = Pt(size)
            p.font.color.rgb = RGBColor.from_string(foreground)
        slide.notes_slide.notes_text_frame.text = 'Replace placeholders with sourced content. Verify every rendered slide.'
    presentation.save(path)


def workbook(name, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    book = Workbook()
    data = book.active
    data.title = 'Inputs'
    summary = book.create_sheet('Summary')
    if name == 'analytics-dashboard':
        data.append(['Metric','Period','Actual','Target','Variance','Variance percent','Source'])
        for row in range(2,12):
            data.cell(row,5,f'=IF(OR(C{row}="",D{row}=""),"",C{row}-D{row})')
            data.cell(row,6,f'=IF(OR(C{row}="",D{row}="",D{row}=0),"",E{row}/D{row})').number_format='0.0%'
        summary.append(['Metric','Actual','Target','Variance'])
        for row in range(2,12):
            for col,source in enumerate(('A','C','D','E'),1):
                summary.cell(row,col,f'=IF(Inputs!A{row}="","",Inputs!{source}{row})')
    elif name == 'financial-budget':
        data.append(['Department','Category','Budget','Actual','Variance','Source'])
        for row in range(2,12):
            data.cell(row,5,f'=IF(OR(C{row}="",D{row}=""),"",D{row}-C{row})')
        summary.append(['Measure','Amount'])
        for row,label,formula in [(2,'Budget','=SUM(Inputs!C2:C11)'),(3,'Actual','=SUM(Inputs!D2:D11)'),(4,'Variance','=B3-B2')]:
            summary.cell(row,1,label); summary.cell(row,2,formula)
    elif name == 'operating-calendar':
        data.append(['Date','Activity','Owner','Status','Notes'])
        summary.append(['Measure','Count'])
        summary.append(['Scheduled activities','=COUNTA(Inputs!B2:B101)'])
        summary.append(['Completed','=COUNTIF(Inputs!D2:D101,"Done")'])
        for row in range(2,12): data.cell(row,1).number_format='yyyy-mm-dd'
    elif name == 'project-tracker':
        data.append(['ID','Task','Owner','Status','Due date','Dependency'])
        summary.append(['Measure','Count'])
        summary.append(['Tasks','=COUNTA(Inputs!A2:A101)'])
        summary.append(['Completed','=COUNTIF(Inputs!D2:D101,"Done")'])
        summary.append(['Blocked','=COUNTIF(Inputs!D2:D101,"Blocked")'])
        for row in range(2,12): data.cell(row,5).number_format='yyyy-mm-dd'
    elif name == 'sales-pipeline':
        data.append(['Opportunity','Account','Stage','Value','Probability','Weighted value','Owner'])
        for row in range(2,12):
            data.cell(row,6,f'=IF(OR(D{row}="",E{row}=""),"",D{row}*E{row})')
            data.cell(row,5).number_format='0%'
        validation=DataValidation(type='decimal',operator='between',formula1=0,formula2=1)
        validation.errorTitle='Probability'; validation.error='Enter a value from 0 to 1.'
        validation.showErrorMessage=True; data.add_data_validation(validation); validation.add('E2:E11')
        summary.append(['Measure','Value'])
        summary.append(['Pipeline','=SUM(Inputs!D2:D11)'])
        summary.append(['Weighted pipeline','=SUM(Inputs!F2:F11)'])
    elif name == 'three-statement-forecast':
        data.append(['Driver','Period 1','Period 2','Period 3'])
        for label in ['Revenue','Direct costs','Operating expenses','Capital spending','Depreciation','Opening cash','Opening fixed assets','Opening debt']:
            data.append([label])
        summary.append(['Statement line','Period 1','Period 2','Period 3'])
        labels=['Revenue','Direct costs','Operating expenses','Depreciation','Net income','Operating cash flow','Capital spending','Closing cash','Fixed assets','Total assets','Debt','Opening equity','Retained earnings','Total liabilities and equity','Balance check']
        for label in labels: summary.append([label])
        for col in range(2,5):
            c=get_column_letter(col); prev=get_column_letter(col-1)
            formulas={2:f'Inputs!{c}2',3:f'Inputs!{c}3',4:f'Inputs!{c}4',5:f'Inputs!{c}6',
                      6:f'{c}2-{c}3-{c}4-{c}5',7:f'{c}6+{c}5',8:f'Inputs!{c}5',
                      9:f'{"Inputs!B7" if col==2 else prev+"9"}+{c}7-{c}8',
                      10:f'{"Inputs!B8" if col==2 else prev+"10"}+{c}8-{c}5',
                      11:f'{c}9+{c}10',12:'Inputs!$B$9',13:'Inputs!$B$7+Inputs!$B$8-Inputs!$B$9',
                      14:f'{c}6' if col==2 else f'{prev}14+{c}6',15:f'{c}12+{c}13+{c}14',16:f'{c}11-{c}15'}
            for row,formula in formulas.items(): summary.cell(row,col,'='+formula)
        notes=book.create_sheet('Scope')
        notes.append(['Simplified example model'])
        notes.append(['Cash sales and cash costs; no tax, working capital, debt changes, or dividends.'])
        notes.append(['Opening balances use Period 1 inputs only. Extend the model for actual requirements.'])
    for sheet in book:
        sheet.freeze_panes='A2'
        sheet.sheet_view.showGridLines=False
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='183B56')
            cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=30
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(name='Arial',size=11,color='183B56')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.data_type=='f' and cell.number_format=='General': cell.number_format='#,##0.00;[Red](#,##0.00);–'
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=24 if col==1 else 18
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'; sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=1
        sheet.print_options.horizontalCentered=True
        sheet.print_area=sheet.dimensions
    book.save(path)


def main():
    with tempfile.TemporaryDirectory(prefix='work-templates-') as temporary:
        staging=Path(temporary)
        for name,sections in DOCUMENTS.items():
            path=staging/f'{name}.docx'; document(name,sections,path)
        for name,sections in DECKS.items():
            path=staging/f'{name}.pptx'; deck(name,sections,path)
        for name in SHEETS:
            path=staging/f'{name}.xlsx'; workbook(name,path)
        for path in sorted(staging.iterdir()):
            description=f'Create a {creator.KINDS[path.suffix][0]} using the original Work {path.stem.replace("-", " ")} reference. Use when this template is selected or requested.'
            result=creator.create(path,path.stem,description,ROOT/'skills')
            skill=Path(result['directory'])/'SKILL.md'
            with skill.open('a') as stream:
                stream.write('\nThis is an original Work reference template. It does not reproduce the OpenAI template design or retained assets.\n')
            print(result['name'])


if __name__=='__main__': main()
